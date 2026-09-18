"""Generate the Word submission document required by the brief.

Run:  backend/.venv/bin/python docs/make_submission_docx.py

The brief requires a single .docx submitted through the applied-job page containing
the repository link, live demo URL, a README overview (architecture, schema, local
setup, CI/CD) and any credentials a reviewer needs.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

GREEN = RGBColor(0x1C, 0x6B, 0x4F)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)

# --- values a candidate must replace before submitting ------------------------
REPO_URL = "https://github.com/<your-username>/darukaa-earth"
DEMO_URL = "https://darukaa-earth.vercel.app"
API_URL = "https://darukaa-api.up.railway.app"
DEMO_EMAIL = "admin@darukaa.earth"
DEMO_PASSWORD = "Admin@12345"
REVIEWER_ACCOUNTS = [
    "ankita.dasgupta@darukaa.com",
    "harsh.kumar@darukaa.com",
    "utkarsh.gauniyal@darukaa.com",
    "guneet.mutreja@darukaa.com",
]

OUTPUT = Path(__file__).with_name("Darukaa_Earth_Submission.docx")


def heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = GREEN if level <= 2 else INK
        run.font.name = "Calibri"


def para(doc: Document, text: str, *, bold: bool = False, muted: bool = False, size: int = 11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if muted:
        run.font.color.rgb = MUTED
    return p


def bullet(doc: Document, text: str, *, level: int = 0) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Pt(18 + level * 18)
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.name = "Calibri"


def code_block(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    shading = p._p.get_or_add_pPr()
    from docx.oxml import OxmlElement

    shd = OxmlElement("w:shd")
    shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", "F1F5F9")
    shading.append(shd)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for index, header in enumerate(headers):
        cell = t.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = "Calibri"
    for row in rows:
        cells = t.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(value)
            run.font.size = Pt(10)
            run.font.name = "Calibri"
    doc.add_paragraph()


def build() -> Document:
    doc = Document()

    # Base style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # --- title ---------------------------------------------------------------
    title = doc.add_heading("Darukaa.Earth", level=0)
    for run in title.runs:
        run.font.color.rgb = GREEN

    para(
        doc,
        "Full-Stack Developer Hackathon — Submission Document",
        bold=True,
        size=13,
    )
    para(
        doc,
        "A geospatial analytics platform for managing and visualising carbon and "
        "biodiversity projects.",
        muted=True,
    )
    doc.add_paragraph()

    # --- 1. links ------------------------------------------------------------
    heading(doc, "1. Submission links", 1)
    table(
        doc,
        ["Item", "Link"],
        [
            ["GitHub repository (private)", REPO_URL],
            ["Live demo URL", DEMO_URL],
            ["Deployed API", f"{API_URL}/api/v1/health"],
            ["API documentation", f"{API_URL}/docs (development builds only)"],
        ],
    )

    heading(doc, "Repository access", 2)
    para(
        doc,
        "The repository is private. Collaborator access has been granted to the "
        "following accounts:",
    )
    for account in REVIEWER_ACCOUNTS:
        bullet(doc, account)
    para(
        doc,
        "If any invitation has not arrived, please let me know and I will re-send it "
        "immediately.",
        muted=True,
    )

    # --- 2. credentials ------------------------------------------------------
    heading(doc, "2. Credentials and review notes", 1)
    para(doc, "A seeded administrator account is available on the live demo:")
    table(
        doc,
        ["Field", "Value"],
        [
            ["Email", DEMO_EMAIL],
            ["Password", DEMO_PASSWORD],
            ["Role", "admin (full read/write)"],
        ],
    )
    para(doc, "What the seeded data contains:")
    bullet(doc, "3 projects across reforestation, agroforestry and biodiversity conservation")
    bullet(doc, "8 sites with drawn polygon boundaries in the Gangetic plain (Delhi NCR and western Uttar Pradesh)")
    bullet(doc, "288 monthly observations spanning 36 months, deterministic and reproducible")
    para(
        doc,
        "Registering a new account also works; note that the first account created on a "
        "fresh database is promoted to admin, and later sign-ups default to read-only "
        "viewer, which is useful for checking the authorisation behaviour.",
        muted=True,
    )
    para(
        doc,
        "Suggested review path: sign in → Dashboard → open a project → open a site to see "
        "the time-series analytics → Site Map → “Add site” to draw a new polygon.",
    )

    # --- 3. architecture -----------------------------------------------------
    heading(doc, "3. README overview", 1)

    heading(doc, "3.1 High-level architecture", 2)
    para(
        doc,
        "Three tiers, each independently deployable. The browser talks only to relative "
        "URLs (/api/v1/...), so the same bundle runs behind a preview proxy, Vercel, or "
        "local development without CORS or host changes.",
    )
    table(
        doc,
        ["Layer", "Technology"],
        [
            ["Frontend", "React 18, TypeScript, Vite, Mapbox GL JS + mapbox-gl-draw, Chart.js"],
            ["State / data", "TanStack Query (caching, invalidation), React Router v6"],
            ["Backend", "Python 3.12, FastAPI, SQLAlchemy 2.0, GeoAlchemy2, Alembic"],
            ["Auth", "JWT access + refresh tokens (PyJWT), bcrypt password hashing"],
            ["Database", "PostgreSQL 17 with PostGIS 3.5"],
            ["CI/CD", "GitHub Actions → Vercel (frontend), Railway (API), Neon (PostGIS)"],
        ],
    )
    para(doc, "Key design decisions:", bold=True)
    bullet(
        doc,
        "Geometry is validated twice on purpose: Shapely validates and repairs incoming "
        "GeoJSON so a hand-drawn self-intersecting polygon becomes a valid MULTIPOLYGON "
        "instead of a 500 error, while PostGIS enforces the column type on write.",
    )
    bullet(
        doc,
        "Area and centroid are computed by PostGIS at write time and stored as columns, "
        "using the geography type for a true geodesic area. On a real drawn parcel this "
        "returned 2,133.52 ha where a naive planar degree-squared conversion gave "
        "2,143.90 ha — a 0.5% error that compounds across a portfolio.",
    )
    bullet(
        doc,
        "Aggregation is pushed into SQL (GROUP BY, date_trunc, COALESCE, DISTINCT ON), so "
        "the Python layer only reshapes rows into chart-ready arrays and response size "
        "stays flat as observation volume grows.",
    )
    bullet(
        doc,
        "The map has an SVG fallback renderer for when no Mapbox token is configured, so "
        "a reviewer without credentials can still see boundaries and exercise the "
        "draw-to-create flow. Both renderers emit identical GeoJSON.",
    )

    # --- 3.2 schema ----------------------------------------------------------
    heading(doc, "3.2 Database schema", 2)
    para(doc, "Four tables, version-controlled with Alembic migrations.")
    code_block(
        doc,
        "users       id UUID PK | email UNIQUE | hashed_password | role(admin|viewer) | is_active\n"
        "projects    id UUID PK | name UNIQUE | project_type | status | country | methodology\n"
        "            | baseline_year | owner_id FK→users ON DELETE CASCADE\n"
        "sites       id UUID PK | name | geometry geometry(MULTIPOLYGON, 4326) + GiST index\n"
        "            | area_hectares | centroid_lat | centroid_lng | land_cover | planting_year\n"
        "            | project_id FK→projects ON DELETE CASCADE | UNIQUE(project_id, name)\n"
        "observations id UUID PK | measured_on | carbon_sequestered_tco2e | biomass_tonnes | ndvi\n"
        "            | canopy_cover_pct | tree_count | biodiversity_index | species_observed\n"
        "            | data_source(satellite|field_survey|modelled)\n"
        "            | site_id FK→sites ON DELETE CASCADE | UNIQUE(site_id, measured_on)",
    )
    para(doc, "Schema rationale:", bold=True)
    bullet(
        doc,
        "MULTIPOLYGON rather than POLYGON: Mapbox GL Draw emits Polygon, which is promoted "
        "on write, so a site can later hold disjoint parcels with no migration.",
    )
    bullet(
        doc,
        "Observations are wide-but-shallow — one row per site per date holding every "
        "metric — rather than an entity-attribute-value table. A site's 36-month history "
        "is a single indexed range scan instead of a pivot join.",
    )
    bullet(
        doc,
        "area_hectares and centroid are denormalised pure functions of the geometry, "
        "recomputed on every geometry write so they cannot drift, removing geodesic maths "
        "from every read path.",
    )
    bullet(
        doc,
        "UNIQUE(site_id, measured_on) makes duplicate measurements impossible at the "
        "database level, not just in application code.",
    )

    # --- 3.3 local setup -----------------------------------------------------
    heading(doc, "3.3 Local setup", 2)
    para(doc, "One command (requires python3, npm, and PostgreSQL with PostGIS):")
    code_block(doc, "./scripts/bootstrap.sh")
    para(doc, "Or with Docker:")
    code_block(
        doc,
        "cp backend/.env.example backend/.env\n"
        "docker compose up --build\n"
        "#  Dashboard  http://localhost:5173\n"
        "#  API docs   http://localhost:8000/docs",
    )
    para(doc, "Manual steps:")
    code_block(
        doc,
        "# database\n"
        "createdb darukaa\n"
        "psql -d darukaa -c \"CREATE EXTENSION IF NOT EXISTS postgis;\"\n\n"
        "# backend\n"
        "cd backend && cp .env.example .env\n"
        "python -m venv .venv && source .venv/bin/activate\n"
        "pip install -r requirements-dev.txt\n"
        "alembic upgrade head\n"
        "python -m app.services.seed\n"
        "uvicorn app.main:app --reload --port 8000\n\n"
        "# frontend (new terminal)\n"
        "cd frontend && npm install && npm run dev",
    )
    para(
        doc,
        "Optional: set VITE_MAPBOX_TOKEN in frontend/.env.local for the satellite basemap. "
        "Without it the app uses the built-in SVG projection and every feature still works.",
        muted=True,
    )

    # --- 3.4 CI/CD -----------------------------------------------------------
    heading(doc, "3.4 CI/CD pipeline", 2)
    para(doc, "Three GitHub Actions workflows:")
    table(
        doc,
        ["Workflow", "Trigger", "What it does"],
        [
            [
                "ci.yml",
                "push / PR to main",
                "Backend job (PostGIS service container): ruff check, ruff format "
                "--check, mypy, alembic upgrade head + alembic check, pytest with coverage. "
                "Frontend job: eslint --max-warnings=0, tsc -b --noEmit, vitest, vite build. "
                "Formatting job: prettier --check across the repo.",
            ],
            [
                "deploy-frontend.yml",
                "main, frontend/** changed",
                "Re-runs lint, typecheck and tests before deploying, builds with the Vercel "
                "CLI and promotes a prebuilt artifact. Writes the deployment URL to the "
                "GitHub environment.",
            ],
            [
                "deploy-backend.yml",
                "main, backend/** changed",
                "Verifies migrations and runs the test suite against a PostGIS service, "
                "deploys via the Railway CLI, then smoke-tests the live /api/v1/health "
                "endpoint with retries so a failed rollout is reported by the pipeline.",
            ],
        ],
    )
    para(
        doc,
        "PostGIS runs as a service container rather than being mocked: the tests exercise "
        "real geometry types, GiST indexes, ST_Area(geography) and Postgres DISTINCT ON, "
        "which no SQLite stand-in can provide.",
    )

    heading(doc, "Code quality enforcement", 2)
    para(
        doc,
        "Husky + lint-staged run on every commit: staged frontend files go through ESLint "
        "--fix and Prettier, staged Python files through ruff format and ruff check --fix, "
        "and the whole tree is then checked with ruff and mypy so results match CI exactly. "
        "A .pre-commit-config.yaml is also included for teams that prefer pre-commit.",
    )
    para(doc, "Current status of every check:")
    code_block(
        doc,
        "ruff check          All checks passed!\n"
        "ruff format         45 files already formatted\n"
        "mypy                Success: no issues found in 37 source files\n"
        "eslint              clean, 0 warnings\n"
        "tsc -b --noEmit     clean\n"
        "pytest              69 passed\n"
        "vitest              42 passed\n"
        "prettier --check    all matched files conform",
    )

    # --- 4. datasets ---------------------------------------------------------
    heading(doc, "4. Datasets and why they were chosen", 1)
    para(
        doc,
        "The brief allows any datasets or mocks provided the choice is documented.",
        muted=True,
    )
    para(
        doc,
        "Site geometries are hand-drawn polygons around real locations in the Gangetic "
        "plain — the Yamuna floodplain near Delhi, agroforestry clusters in Ghaziabad "
        "district, and oxbow wetlands on the upper Ganga — so the map opens on a "
        "recognisable region at realistic parcel sizes.",
    )
    para(
        doc,
        "Observations are synthetic but deterministic, generated from a fixed random seed "
        "with a seasonal NDVI and canopy cycle peaking in the July–August monsoon, a "
        "saturating carbon accumulation curve, and field-survey provenance every sixth "
        "month rather than a single uniform source.",
    )
    para(
        doc,
        "Real satellite archives (Sentinel-2, Landsat, Planet) require API keys and "
        "rate-limited accounts that a reviewer cannot be assumed to have. Deterministic "
        "mocks make the reviewer's experience identical to the author's and keep CI "
        "hermetic — no network, no credentials, no flakiness. Values are calibrated to "
        "plausible published ranges (NDVI 0.3–0.95, canopy 18–98%, 3–8 tCO2e/ha/yr for "
        "Indian tropical plantations) so the charts are meaningful rather than noise.",
    )

    # --- 5. trade-offs -------------------------------------------------------
    heading(doc, "5. Notable trade-offs", 1)
    bullet(
        doc,
        "Sync SQLAlchemy handlers on an async framework — FastAPI runs def endpoints in a "
        "threadpool. SQLAlchemy's async support with GeoAlchemy2 is materially less "
        "mature, and the workload is I/O-bound against a connection pool.",
    )
    bullet(
        doc,
        "JWT rather than server sessions, as specified by the brief. Logout is stateless; "
        "the /auth/logout endpoint exists as the single place to add a denylist later "
        "without changing the client.",
    )
    bullet(
        doc,
        "Tests hit a real PostGIS database (~17s for 69 tests) against a dedicated "
        "<dbname>_test database, so the suite can never destroy a developer's seeded data.",
    )
    bullet(
        doc,
        "mapbox-gl is ~1.9MB, so it is lazy-loaded on the /map route only; the dashboard's "
        "initial payload is ~430KB.",
    )

    doc.add_paragraph()
    closing = doc.add_paragraph()
    run = closing.add_run(
        "Thank you for taking the time to review this submission. I am happy to walk "
        "through any part of the codebase."
    )
    run.italic = True
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    closing.alignment = WD_ALIGN_PARAGRAPH.LEFT

    return doc


def main() -> None:
    doc = build()
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")

    # Verify the file is a readable .docx with the expected sections.
    check = Document(OUTPUT)
    headings = [p.text for p in check.paragraphs if p.style.name.startswith("Heading")]
    print(f"Paragraphs: {len(check.paragraphs)} | Tables: {len(check.tables)}")
    print("Headings:")
    for h in headings:
        print(f"  - {h}")


if __name__ == "__main__":
    main()
